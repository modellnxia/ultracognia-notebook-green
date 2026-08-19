"""
Teste do renderizador PPTX (app/decks/pptx_render.py) — `python-pptx` real,
não mockado, mesmo raciocínio de test_pdf.py: só rodando de verdade (montar
o arquivo e reabrir com a própria lib) prova que o .pptx é válido.
"""

import base64
from unittest.mock import AsyncMock, patch

import pytest
from pptx import Presentation

# PNG 1x1 real e válido (decodável pelo Pillow) — python-pptx usa Pillow
# internamente pra ler dimensões da imagem, não aceita bytes arbitrários.
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)

from app.decks.models import (
    BulletListBlock,
    DeckSpec,
    HeadingBlock,
    LayoutId,
    ParagraphBlock,
    QuoteBlock,
    Slide,
    SlideAsset,
    Theme,
)
from app.decks.pptx_render import render_deck_pptx


def _theme() -> Theme:
    return Theme(
        palette={"primary": "#123456", "background": "#FFFFFF", "text": "#111111"},
        font_stack="Arial, sans-serif",
    )


def _reopen(pptx_bytes: bytes) -> Presentation:
    import io
    return Presentation(io.BytesIO(pptx_bytes))


class TestRenderDeckPptx:
    @pytest.mark.asyncio
    async def test_produces_valid_pptx_with_one_slide_per_layout(self):
        deck = DeckSpec(
            theme=_theme(),
            slides=[
                Slide(layout=LayoutId.COVER, title="Capa"),
                Slide(layout=LayoutId.SECTION_BREAK, title="Seção"),
                Slide(
                    layout=LayoutId.TITLE_BULLETS, title="Pontos",
                    body=[BulletListBlock(items=["Um", "Dois", "Três"])],
                ),
                Slide(
                    layout=LayoutId.TWO_COLUMN, title="Comparação",
                    body=[
                        ParagraphBlock(text="Coluna A"),
                        HeadingBlock(text="Sub"),
                        ParagraphBlock(text="Coluna B"),
                        QuoteBlock(text="Uma citação", attribution="Alguém"),
                    ],
                ),
                Slide(
                    layout=LayoutId.DIAGRAM_FULL, title="Diagrama",
                    body=[ParagraphBlock(text="Legenda")],
                    asset=SlideAsset(kind="diagram", ref="não resolvido"),
                ),
                Slide(layout=LayoutId.CLOSING, title="Obrigado"),
            ],
        )

        pptx_bytes = await render_deck_pptx(deck)

        assert pptx_bytes.startswith(b"PK")  # PPTX é um zip (Open XML)
        prs = _reopen(pptx_bytes)
        assert len(prs.slides._sldIdLst) == 6

    @pytest.mark.asyncio
    async def test_uses_16_9_aspect_ratio(self):
        deck = DeckSpec(theme=_theme(), slides=[Slide(layout=LayoutId.COVER, title="X")])
        pptx_bytes = await render_deck_pptx(deck)
        prs = _reopen(pptx_bytes)
        assert prs.slide_width / prs.slide_height == pytest.approx(16 / 9, rel=0.01)

    @pytest.mark.asyncio
    async def test_titles_and_bullets_appear_as_real_text(self):
        deck = DeckSpec(
            theme=_theme(),
            slides=[
                Slide(
                    layout=LayoutId.TITLE_BULLETS, title="Título Real",
                    body=[BulletListBlock(items=["Item alfa", "Item beta"])],
                ),
            ],
        )
        pptx_bytes = await render_deck_pptx(deck)
        prs = _reopen(pptx_bytes)
        slide = list(prs.slides)[0]
        all_text = "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        assert "Título Real" in all_text
        assert "Item alfa" in all_text
        assert "Item beta" in all_text

    @pytest.mark.asyncio
    async def test_embeds_resolved_image_as_real_picture(self):
        deck = DeckSpec(
            theme=_theme(),
            slides=[
                Slide(
                    layout=LayoutId.DIAGRAM_FULL, title="Com imagem",
                    asset=SlideAsset(kind="image", ref="https://x.supabase.co/sign/foo.png?token=abc"),
                ),
            ],
        )
        with patch(
            "app.decks.pptx_render._fetch_image_bytes",
            new=AsyncMock(return_value=_TINY_PNG),
        ):
            pptx_bytes = await render_deck_pptx(deck)

        prs = _reopen(pptx_bytes)
        slide = list(prs.slides)[0]
        pictures = [sh for sh in slide.shapes if sh.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
        assert len(pictures) == 1

    @pytest.mark.asyncio
    async def test_falls_back_to_placeholder_text_when_asset_unresolved(self):
        deck = DeckSpec(
            theme=_theme(),
            slides=[
                Slide(
                    layout=LayoutId.DIAGRAM_FULL, title="Sem imagem",
                    asset=SlideAsset(kind="diagram", ref="ainda em texto"),
                ),
            ],
        )
        pptx_bytes = await render_deck_pptx(deck)
        prs = _reopen(pptx_bytes)
        slide = list(prs.slides)[0]
        pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
        all_text = "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        assert len(pictures) == 0
        assert "ainda em texto" in all_text

    @pytest.mark.asyncio
    async def test_falls_back_to_placeholder_when_image_download_fails(self):
        deck = DeckSpec(
            theme=_theme(),
            slides=[
                Slide(
                    layout=LayoutId.DIAGRAM_FULL, title="Falha de download",
                    asset=SlideAsset(kind="image", ref="https://x.supabase.co/sign/foo.png?token=abc"),
                ),
            ],
        )
        with patch(
            "app.decks.pptx_render._fetch_image_bytes", new=AsyncMock(return_value=None)
        ):
            pptx_bytes = await render_deck_pptx(deck)

        prs = _reopen(pptx_bytes)
        slide = list(prs.slides)[0]
        pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
        assert len(pictures) == 0  # não quebrou o job, só ficou sem imagem
